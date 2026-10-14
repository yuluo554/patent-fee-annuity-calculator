"""M5：报告导出（XLSX 双格式 + 逐字节可复现 + 内容对等 + 无外链）。

口径（plan/04 §七 / HANDOFF-M5 §二.3）：

* 单一渲染源：Markdown 与 XLSX 消费同一份 ``scan_report_sections`` 章节列表，
  内容对等是**结构性质**——两格式都含同一组章节标题与同一份输入事实表；
* 逐字节可复现：两次导出 sha256 一致（跨进程实测，不靠"应当"）；
* 未启用章节显式写出（K42）；无外链（K42）：两格式正文均不得含 URL。
"""

from __future__ import annotations

import hashlib
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

from conftest import repo_root, data_dir

CLI = [sys.executable, "-X", "utf8", "-m", "patfee.cli"]
SAMPLE = str(repo_root() / "data" / "samples" / "portfolio-demo.json")


def _run_cli(*args: str) -> subprocess.CompletedProcess:
    env = dict(__import__("os").environ)
    env["PYTHONPATH"] = str(repo_root() / "src")
    env.setdefault("PYTHONUTF8", "1")
    return subprocess.run(
        CLI + list(args), capture_output=True, text=True,
        encoding="utf-8", env=env, cwd=str(repo_root()),
    )


def _build_sections():
    """直接经编排层构造章节列表（不走 CLI）。"""
    from decimal import Decimal
    from patfee.rules.loader import load_library
    from patfee.calc import pipeline
    from patfee.calc.portfolio import compare_rounds
    from patfee.cli import _scan_not_enabled_sections
    from patfee.report import scan_report_sections

    payload = json.loads(
        (repo_root() / "data" / "samples" / "portfolio-demo.json").read_text(encoding="utf-8"),
        parse_float=Decimal,
    )
    lib = load_library(data_dir())
    request = pipeline.build_portfolio_request(payload, preview=False)
    result = pipeline.compute_full_portfolio(request, lib)
    baseline, current = pipeline.portfolio_rounds(request)
    reconciliation = compare_rounds(baseline, current)
    return (
        scan_report_sections(
            result, reconciliation, lib, False,
            _scan_not_enabled_sections(), pipeline.DISCLAIMER,
        ),
        result,
    )


def test_scan_report_markdown_structure_preserved(tmp_path):
    """M3 的 Markdown 四块结构在 M5 重构后保持原样（既有断言不放松）。"""
    code = _run_cli("scan", SAMPLE, "--out", str(tmp_path / "unused-out")).returncode
    assert code == 1
    sections, _ = _build_sections()
    from patfee.report.markdown_render import render_markdown

    text = render_markdown(sections)
    assert "# 专利组合费用排期报告" in text
    assert "## 一、输入事实" in text
    assert "## 二、逐项结论与计算过程" in text
    assert "## 三、未启用章节" in text
    assert "## 四、适用依据与免责声明" in text


def test_xlsx_export_reproducible_across_processes(tmp_path):
    """XLSX 逐字节可复现：两次独立进程导出 sha256 一致（D50 顺延项落地）。"""
    out1, out2 = tmp_path / "a", tmp_path / "b"
    r1 = _run_cli("scan", SAMPLE, "--out", str(out1), "--format", "xlsx")
    r2 = _run_cli("scan", SAMPLE, "--out", str(out2), "--format", "xlsx")
    assert r1.returncode == 1 and r2.returncode == 1  # 检出问题 exit 1，不受格式影响
    x1, x2 = out1 / "scan-report.xlsx", out2 / "scan-report.xlsx"
    assert x1.is_file() and x2.is_file()
    h1 = hashlib.sha256(x1.read_bytes()).hexdigest()
    h2 = hashlib.sha256(x2.read_bytes()).hexdigest()
    assert h1 == h2, "两次 XLSX 导出必须逐字节一致（可复现性）"
    # evidence.json 携带 XLSX 指纹且一致
    e1 = json.loads((out1 / "evidence.json").read_text(encoding="utf-8"))
    assert e1["report_xlsx"]["sha256"] == h1
    assert e1["formats"] == ["md", "xlsx"]
    # zip 条目时间戳与存储方式固定
    with zipfile.ZipFile(x1) as z:
        for info in z.infolist():
            assert info.date_time == (1980, 1, 1, 0, 0, 0), info.filename
            assert info.compress_type == zipfile.ZIP_STORED


def test_xlsx_and_markdown_content_equivalence():
    """两格式内容对等（结构性质）：同一份章节列表，标题与表格逐一对上。"""
    sections, result = _build_sections()
    from patfee.report.markdown_render import render_markdown
    from patfee.report.xlsx_render import render_xlsx_bytes
    import io
    from openpyxl import load_workbook

    md = render_markdown(sections)
    wb = load_workbook(io.BytesIO(render_xlsx_bytes(sections)))

    for section in sections:
        # 每个章节在两格式里都有落点：Markdown 标题行 + XLSX 目录表行
        assert section.title in md, section.title
        titles_in_index = [
            row[0] for row in wb["目录"].iter_rows(min_row=2, values_only=True)
        ]
        assert section.title in titles_in_index, section.title

    # 输入事实表逐行对等：XLSX 数据行 = Markdown 表格数据行（行数与首件专利号）
    ws = None
    for name in wb.sheetnames:
        if name.endswith("输入事实"):
            ws = wb[name]
    assert ws is not None
    fact = [s for s in sections if s.title == "一、输入事实"][0]
    xlsx_rows = list(ws.iter_rows(values_only=True))
    assert len(xlsx_rows) >= 1 + len(fact.table_rows())  # 表头 + 数据（+ 表后说明行）
    data_rows = xlsx_rows[1:1 + len(fact.table_rows())]
    # XLSX 空单元格读回 None，Section 里是 None——空即缺，两边一致
    assert [r[1] or "" for r in data_rows] == [
        (row[1] or "") for row in fact.table_rows()
    ]


def test_scan_reports_have_no_external_links(tmp_path):
    """无外链纪律（K42）双格式同守。

    Markdown 正文不得含 URL；XLSX 不得含超链接对象（``<hyperlink>``）与
    外部关系项（``TargetMode="External"``）——OOXML 的 xmlns 规范 URI
    （``http://schemas.openxmlformats.org/...``）是格式固定成分，不是外链。
    """
    out = tmp_path / "r"
    assert _run_cli("scan", SAMPLE, "--out", str(out), "--format", "xlsx").returncode == 1
    md_text = (out / "scan-report.md").read_text(encoding="utf-8")
    assert "http://" not in md_text and "https://" not in md_text
    with zipfile.ZipFile(out / "scan-report.xlsx") as z:
        for name in z.namelist():
            data = z.read(name)
            assert b"<hyperlink" not in data, name
            assert b'TargetMode="External"' not in data, name


def test_xlsx_missing_dependency_is_exit_3_not_crash(tmp_path, monkeypatch):
    """缺 openpyxl：exit 3 + 提示（不抛裸异常）——惰性导入契约与 gui 同构。"""
    from patfee.report.xlsx_render import render_xlsx_bytes, UsageError
    from patfee.report.base import Section

    import builtins
    real_import = builtins.__import__

    def fake_import(name, *a, **kw):
        if name == "openpyxl" or name.startswith("openpyxl."):
            raise ImportError("No module named 'openpyxl' (simulated)")
        return real_import(name, *a, **kw)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    sections, _ = _build_sections()
    with pytest.raises(UsageError) as excinfo:
        render_xlsx_bytes(sections)
    assert excinfo.value.reason_code == "missing_dependency"
    assert "openpyxl" in excinfo.value.message


def test_scan_default_format_stays_markdown(tmp_path):
    """缺省 --format 仍为 md：不写 xlsx，evidence.formats == ['md']（向后兼容）。"""
    out = tmp_path / "r"
    assert _run_cli("scan", SAMPLE, "--out", str(out)).returncode == 1
    assert (out / "scan-report.md").is_file()
    assert not (out / "scan-report.xlsx").exists()
    ev = json.loads((out / "evidence.json").read_text(encoding="utf-8"))
    assert ev["formats"] == ["md"]
